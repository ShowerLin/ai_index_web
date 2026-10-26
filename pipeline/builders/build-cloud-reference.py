"""Extract saved cloud estimates without modifying the Bloomberg workbook."""
import datetime, hashlib, json, math, re
import openpyxl
from paths import DATA_ROOT, BUILD_ROOT

source = DATA_ROOT / 'AI_Monetization.xlsx'
values = openpyxl.load_workbook(source, read_only=True, data_only=True)
formulas = openpyxl.load_workbook(source, read_only=True, data_only=False)
sheet, query = values['Cloud FA Value'], formulas['Cloud FA Formula']
records = []
# Separate native fiscal-period sequences and source scopes are retained.
blocks = [(1, [2, 3, 4]), (6, [7, 8, 9]), (11, [12, 13, 14]),
          (16, [17, 18, 19, 20]), (22, [23, 24, 25]), (26, [27])]
for header, rows in blocks:
    for row in rows:
        label = str(sheet.cell(row, 1).value).strip()
        entity = str(sheet.cell(header, 1).value or 'Oracle').strip()
        if row in (23, 24, 25, 27): entity = label
        metric = 'RPO' if row in (23, 24, 25, 27) else label
        if row == 19: metric = 'Gross profit proxy'
        if row == 20: metric = 'Gross margin proxy'
        for col in range(2, sheet.max_column + 1):
            cell = sheet.cell(row, col)
            value = cell.value
            if value is None: continue
            if not isinstance(value, (int, float)) or not math.isfinite(value):
                raise ValueError(f'Invalid cached value at {cell.coordinate}: {value}')
            period = sheet.cell(header, col).value
            if not re.fullmatch(r'[1-4]Q\d{2}', str(period)):
                raise ValueError(f'Invalid fiscal period: {period}')
            formula = query.cell(row, col).value
            formula = getattr(formula, 'text', formula)
            direct = isinstance(formula, str) and 'BQL(' in formula
            if direct and ('FA_ACT_EST_DATA=E' not in formula or 'CURRENCY=USD' not in formula or '/1M' not in formula):
                raise ValueError(f'Unexpected query basis at {cell.coordinate}')
            margin = 'Margin' in label
            records.append(dict(entity=entity, metric=metric, sourceLabel=label,
                fiscalPeriod=period, value=value if margin else value / 1000,
                unit='fraction' if margin else 'USD billions', estimate=True,
                sourceCell=cell.coordinate, formula=formula,
                basis='Bloomberg estimate' if direct else 'Derived from Bloomberg estimates'))
assert len(records) == sum(len(rows) for _, rows in blocks) * (sheet.max_column - 1), len(records)
payload = dict(metadata=dict(source='AI_Monetization.xlsx', sheet=sheet.title,
    formulaSheet=query.title, workbookSha256=hashlib.sha256(source.read_bytes()).hexdigest(),
    extractedAtUtc=datetime.datetime.now(datetime.timezone.utc).isoformat(),
    sourceRetrievedAt=None, role='Core index reference; no scoring weight',
    periodBasis='Native Bloomberg FPR fiscal labels; calendar alignment not verified',
    limitations=['All queried observations explicitly request estimates (FA_ACT_EST_DATA=E).',
        'Revenue and RPO scopes differ: Google Cloud segment, AWS revenue versus AMZN issuer RPO, Intelligent Cloud versus Microsoft commercial RPO, Oracle cloud/software revenue versus issuer RPO.',
        'Oracle revenue less cost of revenue is a gross-profit proxy, despite the source operating-income label.',
        'RPO is a stock of contracted obligations, not quarterly revenue or cash; no cross-company total or RPO/revenue ratio is calculated.',
        'This sheet contains no CapEx observations. Existing historical CapEx remains the scoring input.']),
    observations=records)
(BUILD_ROOT/'app/cloud-reference-snapshot.ts').write_text(
    'export const cloudReferenceSnapshot = '+json.dumps(payload, indent=2)+' as const;\n', encoding='utf-8')
print('Extracted', len(records), 'cloud estimate reference observations')
