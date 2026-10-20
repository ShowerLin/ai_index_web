"""Deterministic cached normalization with raw-file provenance."""
import calendar, csv, datetime as dt, hashlib, json, re, zipfile
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path
import openpyxl
from bs4 import BeautifulSoup
from .acquire import ROOT

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p, obj): p.write_text(json.dumps(obj, indent=2, ensure_ascii=False, allow_nan=False)+'\n', encoding='utf-8')
def write_csv(p, rows):
    if not rows:
        p.write_text('record_id\n', encoding='utf-8'); return
    fields=list(dict.fromkeys(k for r in rows for k in r))
    with p.open('w', newline='', encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def read_csv(p): return list(csv.DictReader(p.open(encoding='utf-8', newline='')))
def xml_rows(path, sheet=1):
    """Stream cached cells, avoiding styles/charts in the large EIA workbook."""
    with zipfile.ZipFile(path) as z:
        strings=[]
        if 'xl/sharedStrings.xml' in z.namelist():
            for _, node in ET.iterparse(z.open('xl/sharedStrings.xml'), events=['end']):
                if node.tag==NS+'si': strings.append(''.join(x.text or '' for x in node.iter(NS+'t')));node.clear()
        for _, node in ET.iterparse(z.open(f'xl/worksheets/sheet{sheet}.xml'), events=['end']):
            if node.tag!=NS+'row': continue
            values={}
            for c in node.findall(NS+'c'):
                v=c.find(NS+'v')
                if v is None or v.text is None: continue
                val=v.text
                if c.get('t')=='s': val=strings[int(val)]
                elif c.get('t') not in ('str','e'): val=float(val)
                values[re.sub(r'\d','',c.get('r'))]=val
            yield int(node.get('r')),values
            node.clear()

class Build:
    def __init__(self, config, out):
        self.config=config; self.out=Path(out);self.out.mkdir(parents=True,exist_ok=True)
        self.raw=ROOT/'data/raw/electricity';self.run=config['raw_run'];self.tables={};self.sources={};self.reviews=[]
    def path(self, source, name):
        p=self.raw/source/self.run/name
        if not p.is_file(): raise FileNotFoundError(p)
        return p
    def source(self, source, name, vintage, pub=None):
        p=self.path(source,name);m=json.loads(p.with_name(p.name+'.manifest.json').read_text(encoding='utf-8'))
        key=f'{source}:{name}'
        self.sources[key]=dict(m,source_release_id=key,title=name,publisher=source,source_vintage=vintage,
            publication_date=pub,access_method='public_file' if p.suffix!='.json' else 'documented_manual_evidence',
            expected_frequency={'EIA_RETAIL':'monthly/annual','CENSUS':'monthly','CBRE':'half-year','PJM':'annual'}.get(source,'irregular'),
            availability_status='acquired',license_notes='Public research evidence; publisher rights retained',coverage_start=None,coverage_end=None)
        return p,self.sources[key]
    def obs(self, src, locator, metric, value, unit, status='reported_actual', **dims):
        row={k:src[k] for k in ('source_id','source_vintage','publication_date','retrieved_at_utc','run_id','source_url','raw_path')}
        row.update(source_release_id=src['source_release_id'],locator=locator,metric=metric,value=value,unit=unit,
                   status=status,quality_flag='reviewed' if src['publication_date'] else 'publication_date_unknown',
                   notes='' if src['publication_date'] else 'Publication day unavailable; retrieval is not publication.')
        row.update(dims)
        identity=[src['source_release_id'],src['source_vintage'],locator,metric,[(k,dims[k]) for k in sorted(dims) if k not in ('notes','quality_flag')]]
        row['record_id']=hashlib.sha256(json.dumps(identity,sort_keys=True).encode()).hexdigest()[:24]
        return row
    def retail(self):
        rows=[]
        for name,sheet,freq,cols in [('HS861_2010-.xlsx','Total Electric Industry','annual',[3,7,11,15,19]),('HS861M_2010-.xlsx','Monthly-States','monthly',[5,9,13,17,21])]:
            p,src=self.source('EIA_RETAIL',name,'retrieved-2026-10-16')
            s=openpyxl.load_workbook(p,data_only=True,read_only=True)[sheet]
            for i,r in enumerate(s.iter_rows(min_row=4,values_only=True),4):
                if not isinstance(r[0],(int,float)) or r[0]<self.config['historical_start']: continue
                y=int(r[0]);m=int(r[1]) if freq=='monthly' else 1;geo=r[2] if freq=='monthly' else r[1]
                if not isinstance(geo,str):continue
                if len(geo)!=2:continue
                for sector,c in zip(['residential','commercial','industrial','transportation','all'],cols):
                    if not isinstance(r[c],(int,float)):continue
                    start=f'{y}-{m:02d}-01';end=f'{y}-{m:02d}-{calendar.monthrange(y,m)[1]}' if freq=='monthly' else f'{y}-12-31'
                    rows.append(self.obs(src,f'{sheet}!{openpyxl.utils.get_column_letter(c+1)}{i}','retail_sales',float(r[c]),'MWh',
                        period_start=start,period_end=end,frequency=freq,geography_type='country' if geo=='US' else 'state',geography_id=geo,sector=sector,
                        coverage_basis='Retail sales to ultimate customers; 50 states + DC, excluding territories',seasonal_adjustment='NSA',
                        provisional=r[3]!='Final' if freq=='monthly' else None))
        # Monthly workbook separates states and omits a national row. Sum only the
        # complete, mutually exclusive 50-state + DC panel; never include US totals.
        grouped=defaultdict(list)
        for r in rows:
            if r['frequency']=='monthly' and r['geography_type']=='state':grouped[(r['period_start'],r['sector'])].append(r)
        for (period,sector),members in grouped.items():
            if len({r['geography_id'] for r in members})!=51:continue
            if any(r['frequency']=='monthly' and r['geography_id']=='US' and r['period_start']==period and r['sector']==sector for r in rows):continue
            first=members[0];src=self.sources[first['source_release_id']]
            rows.append(self.obs(src,'Monthly-States; sum 50 states + DC for '+period+' / '+sector,'retail_sales',sum(r['value'] for r in members),'MWh',
                period_start=period,period_end=first['period_end'],frequency='monthly',geography_type='country',geography_id='US',sector=sector,
                coverage_basis='Derived sum of complete 50 states + DC retail panel; excludes territories',seasonal_adjustment='NSA',provisional=any(r['provisional'] for r in members)))
        self.tables['electricity_history']=rows
    def census(self):
        rows=[]
        for name,adj in [('privsatime.xlsx','SAAR'),('privtime.xlsx','NSA')]:
            p,src=self.source('CENSUS',name,'October-2026')
            s=openpyxl.load_workbook(p,data_only=True,read_only=True).active
            headers=[c.value for c in s[4]]
            for i,r in enumerate(s.iter_rows(min_row=5,values_only=True),5):
                if not isinstance(r[0],str):continue
                match=re.match(r'([A-Za-z]{3})-(\d{2})',r[0])
                if not match:continue
                date=dt.datetime.strptime(match[0],'%b-%y');period=date.strftime('%Y-%m')
                if date.year<2014:continue
                for label in ('Office','Data center'):
                    c=headers.index(label);v=r[c]
                    if not isinstance(v,(int,float)):continue
                    rows.append(self.obs(src,f'{s.title}!{openpyxl.utils.get_column_letter(c+1)}{i}','construction_spending',v,'USD million',
                       period=period,category_code='census_private_office_data_center' if label=='Data center' else 'census_private_office',
                       category_label=label,ownership='private',seasonal_adjustment=adj,currency='USD',price_basis='nominal',value_usd_million=v,
                       provisional=r[0].endswith(('p','r')),frequency='monthly'))
        for name in ('private.xlsx','privateha2.xlsx'):
            p,src=self.source('CENSUS',name,'October-2026');s=openpyxl.load_workbook(p,data_only=True,read_only=True)['Private'];years=[c.value for c in s[4]]
            for i,r in enumerate(s.iter_rows(values_only=True),1):
                label=str(r[0]).strip()
                if label not in ('Office','Data center3'):continue
                for c,y in enumerate(years):
                    if not str(y).isdigit() or int(y)<2014 or not isinstance(r[c],(int,float)):continue
                    rows.append(self.obs(src,f'Private!{openpyxl.utils.get_column_letter(c+1)}{i}','construction_spending',r[c],'USD million',
                       period=str(y),category_code='census_private_office_data_center' if label=='Data center3' else 'census_private_office',
                       category_label='Data center' if label=='Data center3' else 'Office',ownership='private',seasonal_adjustment='NSA',currency='USD',price_basis='nominal',value_usd_million=r[c],frequency='annual'))
        self.tables['construction_spending']=rows
    def outlook(self):
        rows=[]
        p,src=self.source('EIA_OUTLOOK','STEO_m.xlsx','October-2026')
        s=openpyxl.load_workbook(p,data_only=True,read_only=True)['7atab']
        for rowno,metric in [(21,'total_consumption'),(22,'retail_sales'),(23,'residential_sales'),(24,'commercial_sales'),(25,'industrial_sales'),(26,'transportation_sales'),(27,'direct_use')]:
            r=[c.value for c in s[rowno]]
            for y in range(2022,2028):
                vals=r[2+(y-2022)*12:2+(y-2021)*12]
                if len(vals)!=12 or any(not isinstance(v,(int,float)) for v in vals):continue
                rows.append(self.obs(src,f'7atab!{openpyxl.utils.get_column_letter(3+(y-2022)*12)}{rowno}:{openpyxl.utils.get_column_letter(14+(y-2022)*12)}{rowno}; monthly sum',metric,sum(vals),'TWh',
                   'publisher_forecast' if y>=2026 else 'historical_estimate',target_year=y,geography_id='US',scenario='STEO',coverage_basis='STEO Table 7a; total consumption includes direct use; no DC addition'))
        p,src=self.source('EIA_OUTLOOK','aeotab8.xlsx','AEO-2026-counterfactual-baseline')
        s=openpyxl.load_workbook(p,data_only=True,read_only=True).active
        for rowno,metric in [(61,'residential_sales'),(62,'commercial_sales'),(63,'industrial_sales'),(64,'transportation_sales'),(65,'retail_sales'),(66,'direct_use'),(67,'total_consumption')]:
            for c in range(3, min(s.max_column,29)+1):
                y=2025+c-3;v=s.cell(rowno,c).value
                if not isinstance(v,(int,float)):continue
                rows.append(self.obs(src,f'{s.title}!{openpyxl.utils.get_column_letter(c)}{rowno}',metric,v,'TWh','publisher_forecast',target_year=y,geography_id='US',scenario='Counterfactual Baseline',coverage_basis='AEO2026 Table 8; counterfactual scenario, not a consensus/reference forecast'))
        self.tables['electricity_outlook']=rows
    def benchmarks(self):
        rows=[]
        entries=[('2024.pdf','2024-12','2024-12-19',[(2023,'historical',176,'historical_estimate','printed p5'),(2028,'low',325,'publisher_forecast','printed p6'),(2028,'high',580,'publisher_forecast','printed p6')]),
          ('2025-update.pdf','2025-update-published-2026-06','2026-06-18',[(2024,'historical',192,'historical_estimate','printed p8'),(2028,'reference',464,'publisher_forecast','printed p8'),*[(2030,label,v,'publisher_forecast','Table 2, printed p9') for label,v in [('reference',649),('consolidated_deployment',578),('high_asic',664),('low_accelerator_lifetime',590),('high_inference_energy',782),('compounded_low',521),('compounded_high',843)]]])]
        for name,vintage,pub,values in entries:
            p,src=self.source('LBNL',name,vintage,pub)
            for year,scenario,value,status,locator in values:
                rows.append(self.obs(src,locator,'dc_electricity',value,'TWh',status,target_year=year,geography_id='US',scenario=scenario,
                    methodology='Bottom-up shipment/stock, equipment power and facility overhead model; excludes crypto mining',energy_twh=value,
                    uncertainty_type='publisher_scenario; not confidence interval',notes='2025 update revises history and extends horizon; publication June 2026. Sparse textual/table anchors only.'))
        self.tables['dc_energy']=rows
    def markets(self):
        rows=[]
        data=[('h1-2024','2024-08-19', [('under_construction',3871.8,'MW'),('vacancy',2.8,'%'),('preleased_under_construction',3056.4,'MW')]),
          ('h2-2024','2025-02-26',[('inventory',6922.6,'MW'),('under_construction',6350,'MW'),('vacancy',1.9,'%')]),
          ('h1-2025','2025-09-08',[('inventory',8155,'MW'),('under_construction',5242.5,'MW'),('vacancy',1.6,'%'),('preleased_share',74.3,'%')]),
          ('h1-2026','2026-08-27',[('inventory',10903,'MW'),('under_construction',7481.1,'MW'),('vacancy',1.4,'%'),('preleased_share',80.4,'%'),('net_absorption',1456.2,'MW')])]
        def add(src, vintage, market, universe, country, coverage, values, locator):
            for metric,value,unit in values:
                rows.append(self.obs(src,locator,metric,value,unit,market_id=market,country=country,market_universe_id=universe,
                    as_of=vintage,stage={'inventory':'operating','under_construction':'under_construction','preleased_share':'under_construction_subset','preleased_under_construction':'under_construction_subset'}.get(metric,'not_applicable'),
                    capacity_basis='unknown',coverage_type=coverage,metric_value=value,metric_unit=unit,expected_online_year=None,
                    prelease_denominator='under_construction' if 'prelease' in metric else None,notes='Capacity MW basis not explicitly established in extracted public text; energy conversion blocked. Provider panels overlap and must not be summed.'))
        for vintage,pub,values in data:
            p,src=self.source('CBRE',vintage+'.html',vintage,pub)
            add(src,vintage,'US_primary_8','CBRE_primary_8_US','US','Wholesale colocation: Northern Virginia, Dallas-Fort Worth, Silicon Valley, Chicago, Phoenix, NY Tri-State, Atlanta, Hillsboro',values,'State of the Market bullets')
            if vintage=='h1-2026':
                add(src,vintage,'Northern_Virginia','CBRE_primary_8_US','US','Metro, wholesale colocation',[('inventory',4496.5,'MW'),('under_construction',2420.2,'MW'),('net_absorption',467.6,'MW')],'State of the Market; Northern Virginia bullet')
                add(src,vintage,'Atlanta','CBRE_primary_8_US','US','Metro, wholesale colocation',[('under_construction',2882,'MW')],'State of the Market; Atlanta bullet')
        for name,vintage,pub,values in [('h2-2025.html','h2-2025','2026-02-17',[('under_construction',35,'GW'),('inventory',39,'GW')]),('h1-2026.html','h1-2026','2026-08-11',[('under_construction',66,'GW'),('preleased_share',95,'%'),('vacancy',1,'%'),('net_absorption',25,'GW')])]:
            p,src=self.source('JLL',name,vintage,pub)
            add(src,vintage,'North_America','JLL_expanded_NA','North America','Hyperscale + colocation, expanded frontier coverage; not US total',values,'Public newsroom release; market fundamentals')
            if vintage=='h2-2025':add(src,vintage,'Texas','JLL_state_US','US','State; expanded hyperscale and colocation',[('under_construction',6.5,'GW')],'Texas leads frontier market surge')
            else:
                add(src,vintage,'Texas','JLL_state_US','US','State; operating plus construction combined',[('operating_plus_under_construction',26,'GW')],'Public release; Texas comparison')
                add(src,vintage,'Virginia','JLL_state_US','US','State; operating plus construction combined',[('operating_plus_under_construction',13,'GW')],'Public release; Virginia comparison')
                add(src,vintage,'Phoenix','JLL_named_markets','US','Metro, publisher coverage',[('under_construction',1.7,'GW'),('net_absorption',442,'MW')],'Public release; named market observations')
                add(src,vintage,'Atlanta','JLL_named_markets','US','Metro, publisher coverage',[('net_absorption',918,'MW')],'Public release; named market observations')
        self.tables['market_pipeline']=rows
    def pjm(self):
        p,src=self.source('PJM','reviewed-summary.json','January-2026','2026-01-14')
        evidence=json.loads(p.read_text(encoding='utf-8'));rows=[]
        for r in evidence['rows']:
            y,summer,change,winter,wchange=r
            for season,v,delta in [('summer',summer,change),('winter',winter,wchange)]:
                for vintage,value in [('2026',v),('2025',v-delta)]:
                    rows.append(self.obs(dict(src,source_vintage=vintage),f'Changes From 2025 Long-Term Load Forecast table, {y}, {season}',f'{season}_peak',value,'MW','publisher_forecast',
                        target_year=y,zone_id='PJM_RTO',geography_id='PJM_RTO',season=season,weather_basis='normal forecast; unrestricted',
                        load_adjustment_basis='Includes large-load/data-center, EV and BTM adjustments; before load management',metric_unit='MW',dc_component_disclosed=False,
                        notes='2025 value reconstructed exactly as 2026 published peak minus published revision. Winter year denotes season ending in target year. No zonal summation.'))
        self.tables['pjm_forecasts']=rows
    def hourly(self):
        p,src=self.source('EIA_GRID','PJM.xlsx','retrieved-2026-10-16');rows=[];epoch=dt.datetime(1899,12,30)
        for i,r in xml_rows(p):
            if i==1:
                assert r['B']=='UTC time' and r['I']=='Demand';continue
            if not isinstance(r.get('B'),(int,float)):continue
            stamp=(epoch+dt.timedelta(days=r['B'])).replace(microsecond=0)
            # Source timestamps are hour-ending. Whole-hour rounding removes Excel fraction noise.
            stamp=stamp.replace(minute=0,second=0)+(dt.timedelta(hours=1) if stamp.minute>30 else dt.timedelta())
            if stamp.date()>dt.date.fromisoformat(self.config['cutoff_date']):continue
            value=r.get('I') if isinstance(r.get('I'),(int,float)) else None
            rows.append(self.obs(src,f'Published Hourly Data!I{i}','demand',value,'MW',timestamp_utc=stamp.isoformat()+'Z',source_timezone='UTC (hour-ending); Eastern local columns unused',
                interval_hours=1,geography_id='PJM',demand_mw=value,missing_reason='source_blank' if value is None else '',provisional=True))
        self.tables['hourly_load']=rows
    def build(self):
        for method in ('retail','census','outlook','benchmarks','markets','pjm','hourly'):
            getattr(self,method)();print(method,'normalized',flush=True)
        for name,rows in self.tables.items():write_csv(self.out/(name+'.csv'),rows)
        for meta in sorted(self.raw.glob('*/'+self.run+'/*.manifest.json')):
            m=json.loads(meta.read_text(encoding='utf-8'));name=meta.name.removesuffix('.manifest.json');key=m['source_id']+':'+name
            if key not in self.sources:
                self.sources[key]=dict(m,source_release_id=key,title=name,publisher=m['source_id'],source_vintage='captured-2026-10-16',publication_date=None,
                    access_method='public_file',expected_frequency='irregular',availability_status='archived_not_normalized',license_notes='Publisher rights retained',coverage_start=None,coverage_end=None)
        write_csv(self.out/'sources.csv',list(self.sources.values()))
        dump(self.out/'sources.json',list(self.sources.values()))
        gapfile=ROOT/'data/runs'/self.run/'electricity/acquisition_gaps.json'
        gaps=json.loads(gapfile.read_text(encoding='utf-8')) if gapfile.exists() else []
        for source,field,reason in [
          ('EIA_GRID','national hourly panel','PJM acquired; US48 workbook endpoint returned HTML. National BA panel not assembled.'),
          ('EIA_OUTLOOK','archived STEO vintages','Attempted October 2025 XLSX URL returned 404; current STEO and AEO tables acquired.'),
          ('PJM','2024/2025/2026 workbooks, zonal energy, midyear and accuracy tables','Direct requests returned 403/SSL errors. Public 2026 summary provides fixed-target peaks and exact 2025 revisions; third annual vintage and zonal tables unavailable.'),
          ('CBRE','full market/vintage panel and IT basis','Public summaries acquired for four editions; embedded chart tables not extracted; H2 2025 URL 404. Unknown power basis blocks energy conversion.'),
          ('JLL','US national capacity history','North America expanded coverage includes Canada and hyperscale; US named market/state facts stored separately. No national allocation inferred.'),
          ('MODEL','observed bottom-up cohort and backtest','No verified IT MW plus delivery-date cohort available. Illustrative cohort stress tests only; no real-time backtest claimed.'),
          ('CENSUS','Bloomberg comparable periods','CNSTPRDA workbook is filled observation dates with unverified period mapping. Source-specific comparison withheld; cached Bloomberg unchanged.')]:
            gaps.append(dict(source_id=source,requested_field=field,reason=reason,impact='Partial research coverage',attempted_url='',attempted_at='2026-10-16',fallback='Explicit scope limits; no imputation',resolution_status='open'))
        write_csv(self.out/'gaps.csv',gaps)
        dump(self.out/'build_manifest.json',dict(schema_version=1,cutoff=self.config['cutoff_date'],config=self.config,table_counts={k:len(v) for k,v in self.tables.items()},
            input_hashes={s['raw_path']:s['checksum_sha256'] for s in self.sources.values()},coverage_status='partial',core_index_changed=False))
        return self.out
