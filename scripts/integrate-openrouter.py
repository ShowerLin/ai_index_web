"""Read cached workbook values; regenerate the embedded research snapshot without editing Excel."""
import calendar, collections, datetime as dt, hashlib, json, pathlib, re, sys
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
SOURCE = pathlib.Path(sys.argv[1]) if len(sys.argv)>1 else pathlib.Path('/Users/linyu/Library/CloudStorage/OneDrive-Personal/AI_INDEX/Indexlist.xlsx')

def extract(source):
    w = openpyxl.load_workbook(source, read_only=True, data_only=True)
    rows = list(w['OpenRouter Token Value'].values)
    buckets = {}; seen = set(); daily = collections.defaultdict(int)
    for row in rows[1:]:
        if not row or not isinstance(row[0], dt.datetime): continue
        date, model, author, tokens = row[:4]
        assert isinstance(tokens, (int,float)) and tokens >= 0
        key = (date,model); assert key not in seen, f'Duplicate {key}'; seen.add(key)
        month = date.strftime('%Y-%m'); daily[date.date()] += int(tokens)
        b = buckets.setdefault(month,dict(month=month,total=0,days=set(),authors=collections.defaultdict(int),free=0,other=0))
        b['total'] += int(tokens); b['days'].add(date.day); b['authors'][author] += int(tokens)
        if ':free' in model: b['free'] += int(tokens)
        if model == 'other': b['other'] += int(tokens)
    months = []
    for key,b in sorted(buckets.items()):
        y,m = map(int,key.split('-')); expected = calendar.monthrange(y,m)[1]
        b['missingDays'] = sorted(set(range(1,expected+1))-b['days'])
        b['days'] = len(b['days']); b['expected'] = expected; b['complete'] = b['days']==expected
        months.append(b)
    # Independently reconcile every stored monthly total against daily source records.
    summary = {r[5].strftime('%Y-%m'): r[6] for r in rows[1:] if len(r)>6 and isinstance(r[5],dt.datetime)}
    assert all(b['total']==summary[b['month']] for b in months)
    financial = {}; fr = list(w['Hyperscaler Core Data Value'].values)
    for start in [1,12,23,34,45]:
        ticker = fr[0][start].split()[0]; series=[]
        for rn,row in enumerate(fr[2:],3):
            if not row or not isinstance(row[0],dt.datetime): continue
            vals = {fr[1][c].lstrip('#'): (row[c]/1e9 if c<len(row) and isinstance(row[c],(int,float)) else None) for c in range(start,start+11)}
            if vals['NOPAT'] is None: continue
            date=row[0]; vals.update(date=date.strftime('%Y-%m-%d'),quarter=f'{date.year}Q{(date.month-1)//3+1}',sourceRange=f'{openpyxl.utils.get_column_letter(start+1)}{rn}:{openpyxl.utils.get_column_letter(start+11)}{rn}')
            vals['cash_capex'] = -vals['cash_capex'] if vals['cash_capex'] is not None else None
            vals['ic'] = vals['shareholders_equity']+vals['total_debt']-vals['cash_and_short_term_investments']
            series.append(vals)
        financial[ticker]=series
    missing=[]; d=min(daily)
    while d<=max(daily):
        if d not in daily: missing.append(d.isoformat())
        d+=dt.timedelta(days=1)
    return dict(months=months,financial=financial,missingDays=missing,rows=len(seen),sha256=hashlib.sha256(source.read_bytes()).hexdigest(),asOf='2026-09-10T09:08:08.030Z',through=max(daily).isoformat())

if __name__=='__main__':
    data=extract(SOURCE)
    html=(ROOT/'index.html').read_text()
    html=re.sub(r'<!-- OPENROUTER START -->.*?<!-- OPENROUTER END -->','',html,flags=re.S)
    css=(ROOT/'research.css').read_text(); js=(ROOT/'research.js').read_text()
    block='<!-- OPENROUTER START --><style>'+css+'</style><section id="demand-payback" class="section research"><div class="eyebrow">BUILDOUT → MONETIZATION → PAYBACK</div><h2>Is usage catching up with investment?</h2><p class="research-intro">OpenRouter activity alongside hyperscaler spending, profit and consolidated capital returns.</p><div id="research-app"></div></section><script id="research-data" type="application/json">'+json.dumps(data,separators=(',',':'))+'</script><script>'+js+'</script><!-- OPENROUTER END -->'
    assert '<section class="section methodology"' in html
    html=html.replace('<section class="section methodology"',block+'<section class="section methodology"',1)
    if 'href="#demand-payback"' not in html: html=html.replace('</nav>','<a href="#demand-payback">Demand &amp; Payback</a></nav>',1)
    html=html.replace('Updated · 25 Aug 2026','Core snapshot · 25 Aug 2026')
    (ROOT/'index.html').write_text(html)
    (ROOT/'research-snapshot.json').write_text(json.dumps(data,indent=2))
    print(json.dumps(dict(months=len(data['months']),rows=data['rows'],missing=data['missingDays'],financial={k:len(v) for k,v in data['financial'].items()})))
