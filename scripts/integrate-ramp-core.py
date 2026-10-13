"""Add paid adoption and nominal spending intensity to methodology v2.2.

Run after update-demand-core.py. Repeated runs use the preserved v2.2 input.
"""
import argparse
import copy
import datetime
import hashlib
import json
import statistics
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT = Path('/Users/linyu/Library/CloudStorage/OneDrive-Personal/AI_INDEX/Indexlist.xlsx')

def read_snapshot(path):
    text = path.read_text()
    return json.loads(text[text.index('{'):text.rindex(' as const')])

def shift(month, offset):
    total = int(month[:4]) * 12 + int(month[5:]) - 1 + offset
    return f'{total // 12:04}-{total % 12 + 1:02}'

def integrate(base, monthly, workbook_hash):
    result = copy.deepcopy(base)
    adoption = {r['month']:r['paidAdoption'] for r in monthly if r['paidAdoption'] is not None}
    # July-August 2023 are Filter ALL, not the dashboard Overall series. Exclude that splice from scoring.
    spending = {r['month']:r['medianSpend'] for r in monthly if r['month'] >= '2023-09' and r['medianSpend'] is not None}
    intensity = {}
    for month in spending:
        current = [spending.get(shift(month,-i)) for i in range(3)]
        previous = [spending.get(shift(month,-12-i)) for i in range(3)]
        if all(v is not None for v in current+previous) and sum(previous)>0:
            intensity[month] = (sum(current)/sum(previous)-1)*100
    end = base['metadata']['baselineEnd']
    def stats(series):
        values = [v for m,v in sorted(series.items()) if m<=end]
        return statistics.mean(values), statistics.stdev(values), len(values)
    a_stats, s_stats = stats(adoption), stats(intensity)
    def row(month,ticker,name,pillar,series,level,calibration,treatment,weight,source_field):
        mean,std,count = calibration
        signal = series.get(month)
        score = None if signal is None else max(0,min(100,50+15*(signal-mean)/std))
        previous_level = (adoption if pillar=='Adoption' else spending).get(shift(month,-1))
        momentum = None if level is None or previous_level is None else (level-previous_level if pillar=='Adoption' else (level/previous_level-1)*100 if previous_level>0 else None)
        return dict(ticker=ticker,name=name,pillar=pillar,weight=weight,growth=signal,score=score,
                    contribution=None if score is None else weight*(score-50),baselineCount=count,
                    treatment=treatment,level=level,shortMomentum=momentum,baselineMean=mean,baselineStd=std,
                    source='S14',sourceField=source_field)
    for month in result['months']:
        m = month['month']
        for r in month['rows']:
            if r['pillar']=='Adoption':r['weight']=.2/3
            elif r['pillar']=='Demand':r['weight']=.2/4
            r['contribution']=None if r['score'] is None else r['weight']*(r['score']-50)
        month['rows'].insert(2,row(m,'RAMP_PAID_ADOPTION','Ramp paid business AI adoption','Adoption',adoption,adoption.get(m),a_stats,'Paid adoption level (%)',.2/3,'adoption_rate_pct'))
        insert = max(i for i,r in enumerate(month['rows']) if r['pillar']=='Demand')+1
        month['rows'].insert(insert,row(m,'RAMP_MEDIAN_SPEND','Ramp median AI spend per employee','Demand',intensity,spending.get(m),s_stats,'3M-average median spend YoY growth (%)',.2/4,'median_pepm'))
        month['score']=sum(r['weight']*r['score'] for r in month['rows']) if all(r['score'] is not None for r in month['rows']) else None
    result['metadata'].update(version='2.3',previousVersion='2.2',rampWorkbookSha256=workbook_hash,
        rampSource='Indexlist.xlsx / Ramp AI Data, public dashboard snapshot retrieved 13 October 2026',
        rampMethod='Ramp paid adoption level is one third of Adoption (6.6667% of core). Nominal median monthly spend per employee: YoY growth in trailing 3M average is one quarter of Demand (5% of core). Fixed sample-standard-deviation baselines through June 2026. Spending uses dashboard Overall from September 2023; earlier Filter ALL values excluded. Overlapping provider/sector cohorts and AI spend share are context only.',
        rampBaseline={'adoption':dict(mean=a_stats[0],std=a_stats[1],count=a_stats[2]),'spend':dict(mean=s_stats[0],std=s_stats[1],count=s_stats[2])})
    result['metadata']['previousDemandMethod']=base['metadata']['demandMethod']
    result['metadata']['demandMethod']='DRAM, NAND and TSMC TTM YoY retain their fixed v2.2 baselines. Ramp median spending uses 3M-average YoY. Each receives 5% of the core. TSMC pre-July-2024 levels retain official backfills and a 19-observation baseline. Ramp spending has a 20-observation baseline.'
    return result, intensity

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--input',type=Path,default=DEFAULT);args=parser.parse_args()
    workbook=openpyxl.load_workbook(args.input,read_only=True,data_only=True)
    sheet=workbook['Ramp AI Data'];monthly=[]
    for cells in sheet.iter_rows(min_row=10,min_col=1,max_col=12,values_only=True):
        if not isinstance(cells[0],datetime.datetime):break
        monthly.append(dict(month=cells[0].strftime('%Y-%m'),paidAdoption=None if cells[1] is None else round(cells[1]*100,8),
                            medianSpend=cells[4],spendShare=cells[9],adoptionMom=cells[2],adoptionYoy=cells[3]))
    # Only default dashboard observations for supporting vendor / sector breakdowns.
    breakdown=[]
    for cells in sheet.iter_rows(min_row=10,max_row=sheet.max_row,min_col=14,max_col=23,values_only=True):
        if cells[1]=='Paid adoption' and cells[4]=='Paid adoption share' and cells[8]=='R1':
            breakdown.append(dict(month=cells[0].strftime('%Y-%m'),dimension=cells[2],cohort=cells[3],share=round(cells[5]*100,8)))
    workbook.close()
    active=read_snapshot(ROOT/'app/core-contribution-snapshot.ts')
    preserved=ROOT/'data/core-v2.2.ts'
    if active['metadata']['version']=='2.2':
        preserved.write_text('export const coreContributionSnapshot = '+json.dumps(active,indent=2)+' as const;\n')
    base=read_snapshot(preserved)
    sha=hashlib.sha256(args.input.read_bytes()).hexdigest()
    result,intensity=integrate(base,monthly,sha)
    ramp=dict(metadata=dict(source='https://ramp.com/data/ai-index',retrieved='2026-10-13',through=monthly[-1]['month'],workbookSha256=sha,
             access='Public dashboard payload; undocumented website JSON filters. Documented API requires credentials.'),
             monthly=[dict(r,spendSignal=intensity.get(r['month'])) for r in monthly],breakdown=breakdown)
    (ROOT/'app/ramp-snapshot.ts').write_text('export const rampSnapshot = '+json.dumps(ramp,indent=2)+' as const;\n')
    (ROOT/'app/core-contribution-snapshot.ts').write_text('export const coreContributionSnapshot = '+json.dumps(result,indent=2)+' as const;\n')
    latest=next(m for m in reversed(result['months']) if m['score'] is not None)
    old=next(m for m in base['months'] if m['month']==latest['month'])
    print(json.dumps(dict(month=latest['month'],previous=old['score'],score=latest['score'],delta=latest['score']-old['score'],rampBaseline=result['metadata']['rampBaseline'])))

if __name__=='__main__':main()
